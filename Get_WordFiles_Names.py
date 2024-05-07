import os
from tkinter import Tk, Button, Label, filedialog
from docx import Document
from openpyxl import Workbook

def get_word_files_info(folder_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["File Name", "Page Count"])

    for filename in os.listdir(folder_path):
        if filename.endswith('.docx'):
            docx_path = os.path.join(folder_path, filename)
            doc = Document(docx_path)
            page_count = count_pages(doc)
            ws.append([filename, page_count])

    excel_file_path = os.path.join(folder_path, 'word_files_info.xlsx')
    wb.save(excel_file_path)
    print("Excel file created successfully with word files information.")

def count_pages(doc):
    page_count = 0
    for paragraph in doc.paragraphs:
        page_count += paragraph.text.count('\x0c')
    return page_count

def select_folder():
    folder_path = filedialog.askdirectory()
    if folder_path:
        get_word_files_info(folder_path)
        status_label.config(text="Excel file created successfully with word files information.")

# Create Tkinter window
root = Tk()
root.title("Word Files Info Extractor")

# Create label
instruction_label = Label(root, text="Please select a folder containing Word files:")
instruction_label.pack(pady=10)

# Create button to select folder
select_button = Button(root, text="Select Folder", command=select_folder)
select_button.pack(pady=5)

# Create status label
status_label = Label(root, text="")
status_label.pack(pady=5)

# Run Tkinter event loop
root.mainloop()
